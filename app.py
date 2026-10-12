# app.py
import io
import os
import streamlit as st
import pypdf
from docx import Document
from google import genai
from google.genai import types

# Page Configuration
st.set_page_config(
    page_title="AcademicWrite AI",
    page_icon="🎓",
    layout="wide",
    initial_sidebar_state="expanded",
)

# Custom Styling (Clean Blue-and-White Academic Theme)
st.markdown(
    """
    <style>
    .main {
        background-color: #f8fafc;
    }
    .sidebar .sidebar-content {
        background-color: #ffffff;
    }
    h1, h2, h3 {
        color: #1e3a8a !important;
        font-family: 'Helvetica Neue', Helvetica, Arial, sans-serif;
    }
    .stButton>button {
        background-color: #2563eb;
        color: white;
        border-radius: 6px;
        padding: 0.5rem 1rem;
        border: none;
        font-weight: 600;
    }
    .stButton>button:hover {
        background-color: #1d4ed8;
    }
    .stDownloadButton>button {
        background-color: #059669;
        color: white;
        border-radius: 6px;
        padding: 0.5rem 1rem;
        border: none;
        font-weight: 600;
    }
    .stDownloadButton>button:hover {
        background-color: #047857;
    }
    div.stAlert {
        border-radius: 6px;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

# Sidebar Configuration
with st.sidebar:
    st.image(
        "https://img.icons8.com/color/96/graduation-cap.png", width=64
    )
    st.title("AcademicWrite AI")
    st.markdown(
        "*Your AI-powered thesis and laboratory report writing assistant.*"
    )

    st.markdown("---")
    st.subheader("Configuration")

    # Document Type Selection
    doc_type = st.selectbox(
        "Document Type", ["Laboratory Report", "Thesis Section"]
    )

    # Citation Style Selection
    citation_style = st.selectbox("Citation Style", ["APA", "IEEE", "MLA"])

    st.markdown("---")
    st.subheader("Gemini API Configuration")

    # API Key Handling (Streamlit secrets or Sidebar Password input)
    api_key = None
    if (
        hasattr(st, "secrets")
        and "GEMINI_API_KEY" in st.secrets
        and st.secrets["GEMINI_API_KEY"]
    ):
        api_key = st.secrets["GEMINI_API_KEY"]
        st.success("API Key loaded from secrets.")
    else:
        api_key = st.text_input(
            "Gemini API Key", type="password", help="Enter your Google Gemini API key."
        )
        if not api_key:
            st.warning("Please provide a Gemini API key to proceed.")

    st.markdown("---")
    st.markdown(
        "**About**\n\nBuilt for college students to assist in drafting structured academic sections, refining text, and formatting citations responsibly."
    )


# Helper function to initialize Gemini client safely
def get_gemini_client(key):
    if not key:
        return None
    try:
        return genai.Client(api_key=key)
    except Exception as e:
        return None


# Helper function to generate DOCX file in memory
def create_docx(title, content, citation_info=None):
    doc = Document()
    doc.add_heading(title, level=1)

    # Add paragraphs from content
    paragraphs = content.split("\n")
    for p in paragraphs:
        if p.strip():
            doc.add_paragraph(p.strip())

    if citation_info:
        doc.add_heading("References / Citations", level=2)
        for cit in citation_info:
            doc.add_paragraph(cit)

    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio


# Main Application Tabs
tab1, tab2, tab3 = st.tabs(
    ["📝 Document Builder", "✨ Grammar & Tone Refiner", "📚 Citation Formatter"]
)

# ---------------------------------------------------------
# TAB 1: Document Builder
# ---------------------------------------------------------
with tab1:
    st.header("Document Builder")
    st.write(
        "Generate structured sections for your thesis or laboratory report based on your raw notes and findings."
    )

    col1, col2 = st.columns(2)
    with col1:
        report_title = st.text_input(
            "Report Title or Research Topic",
            placeholder="e.g., Analysis of Enzyme Kinetics at Varying Temperatures",
        )
        section_type = st.selectbox(
            "Select Section",
            [
                "Abstract",
                "Introduction",
                "Objectives",
                "Methodology",
                "Results",
                "Discussion",
                "Conclusion",
            ],
        )

    with col2:
        writing_tone = st.selectbox(
            "Academic Writing Tone",
            ["Formal & Objective", "Analytical", "Concise", "Descriptive"],
        )
        output_length = st.selectbox(
            "Approximate Output Length", ["Short (~150 words)", "Medium (~300 words)", "Detailed (~500 words)"]
        )

    user_notes = st.text_area(
        "Experiment Details, Notes, Observations, or Raw Data",
        placeholder="Enter your notes, measurements, observations, or outline here...",
        height=150,
    )

    if "builder_output" not in st.session_state:
        st.session_state.builder_output = ""

    if st.button("Generate Content", key="generate_builder"):
        if not api_key:
            st.error("Please enter a valid Gemini API Key in the sidebar or secrets.")
        elif not report_title.strip() or not user_notes.strip():
            st.warning("Please provide both a report title/topic and your experiment notes/content.")
        else:
            client = get_gemini_client(api_key)
            if not client:
                st.error("Failed to initialize Gemini client. Check your API key.")
            else:
                with st.spinner("Generating academic content using Gemini..."):
                    try:
                        # Construct prompt based on doc_type and strict academic constraints
                        constraint_instruction = ""
                        if doc_type == "Laboratory Report":
                            constraint_instruction = (
                                "CRITICAL CONSTRAINT FOR LAB REPORTS: Preserve the user's actual experiment details "
                                "and results strictly. NEVER invent experimental measurements or fabricate data."
                            )
                        else:
                            constraint_instruction = (
                                "CRITICAL CONSTRAINT FOR THESIS SECTIONS: Distinguish supplied facts from any assumptions "
                                "and clearly note or ask for missing information if needed."
                            )

                        prompt = f"""
                        You are an expert academic writing assistant. Write a professional {section_type} for a {doc_type} titled "{report_title}".
                        Tone: {writing_tone}.
                        Target Length: {output_length}.
                        
                        Supplied Notes/Data:
                        {user_notes}
                        
                        {constraint_instruction}
                        
                        Produce clear, rigorous academic prose suitable for submission.
                        """

                        response = client.models.generate_content(
                            model="gemini-3.8-flash",
                            contents=prompt,
                        )
                        st.session_state.builder_output = response.text
                        st.success("Content generated successfully!")
                    except Exception as e:
                        st.error(f"An error occurred while calling the Gemini API: {e}")

    # Editable Output Field
    if st.session_state.builder_output:
        st.subheader("Generated Content (Editable)")
        st.session_state.builder_output = st.text_area(
            "Edit generated content below before downloading:",
            value=st.session_state.builder_output,
            height=250,
        )

        col_b1, col_b2 = st.columns(2)
        with col_b1:
            if st.button("Improve Generated Content", key="improve_builder"):
                if not api_key:
                    st.error("API Key required.")
                else:
                    client = get_gemini_client(api_key)
                    with st.spinner("Refining content..."):
                        try:
                            refine_prompt = f"""
                            Improve the following academic text for clarity, scholarly tone, and grammar while keeping all factual information intact:
                            
                            {st.session_state.builder_output}
                            """
                            response = client.models.generate_content(
                                model="gemini-3.8-flash",
                                contents=refine_prompt,
                            )
                            st.session_state.builder_output = response.text
                            st.rerun()
                        except Exception as e:
                            st.error(f"Error during improvement: {e}")

        with col_b2:
            docx_file = create_docx(f"{report_title} - {section_type}", st.session_state.builder_output)
            st.download_button(
                label="Download as DOCX",
                data=docx_file,
                file_name=f"{section_type.lower().replace(' ', '_')}_report.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )


# ---------------------------------------------------------
# TAB 2: Grammar & Tone Refiner
# ---------------------------------------------------------
with tab2:
    st.header("Grammar & Tone Refiner")
    st.write(
        "Polish existing draft text, correct grammatical errors, and elevate academic tone without altering your facts."
    )

    refine_mode = st.radio(
        "Input Method", ["Paste Text", "Upload Document (TXT, DOCX, PDF)"], horizontal=True
    )

    raw_text = ""
    if refine_mode == "Paste Text":
        raw_text = st.text_area(
            "Paste your text here",
            placeholder="Paste text from your draft...",
            height=150,
        )
    else:
        uploaded_file = st.file_uploader(
            "Upload Document", type=["txt", "docx", "pdf"]
        )
        if uploaded_file is not None:
            file_extension = uploaded_file.name.split(".")[-1].lower()
            try:
                if file_extension == "txt":
                    raw_text = uploaded_file.getvalue().decode("utf-8")
                elif file_extension == "docx":
                    doc = Document(uploaded_file)
                    raw_text = "\n".join([p.text for p in doc.paragraphs])
                elif file_extension == "pdf":
                    reader = pypdf.PdfReader(uploaded_file)
                    raw_text = "\n".join(
                        [page.extract_text() for page in reader.pages if page.extract_text()]
                    )
                st.info(f"Successfully loaded {uploaded_file.name} ({len(raw_text)} characters).")
            except Exception as e:
                st.error(f"Error reading uploaded file: {e}")

    refinement_option = st.selectbox(
        "Refinement Option",
        [
            "Grammar Correction",
            "Academic Tone",
            "Clarity Improvement",
            "Conciseness",
        ],
    )

    if "refiner_output" not in st.session_state:
        st.session_state.refiner_output = ""

    if st.button("Refine Text", key="run_refiner"):
        if not api_key:
            st.error("Please enter a valid Gemini API Key in the sidebar or secrets.")
        elif not raw_text.strip():
            st.warning("Please provide or upload text to refine.")
        else:
            client = get_gemini_client(api_key)
            if not client:
                st.error("Failed to initialize Gemini client.")
            else:
                with st.spinner("Refining text with Gemini..."):
                    try:
                        prompt = f"""
                        You are an expert academic editor. Apply the following refinement option: '{refinement_option}' to the text below.
                        CRITICAL RULES:
                        1. Preserve original meaning and factual information completely.
                        2. Do not add unsupported facts, fabricated results, or fake references.
                        
                        Original Text:
                        {raw_text}
                        """
                        response = client.models.generate_content(
                            model="gemini-3.8-flash",
                            contents=prompt,
                        )
                        st.session_state.refiner_output = response.text
                        st.success("Text refined successfully!")
                    except Exception as e:
                        st.error(f"An error occurred: {e}")

    if raw_text and st.session_state.refiner_output:
        col_r1, col_r2 = st.columns(2)
        with col_r1:
            st.subheader("Original Text")
            st.text_area("Original", value=raw_text, height=250, disabled=True, key="orig_box")
        with col_r2:
            st.subheader("Refined Text (Editable)")
            st.session_state.refiner_output = st.text_area(
                "Refined", value=st.session_state.refiner_output, height=250, key="ref_box"
            )

        docx_refined = create_docx("Refined Academic Text", st.session_state.refiner_output)
        st.download_button(
            label="Download Refined Document as DOCX",
            data=docx_refined,
            file_name="refined_academic_text.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )


# ---------------------------------------------------------
# TAB 3: Citation Formatter
# ---------------------------------------------------------
with tab3:
    st.header("Citation Formatter")
    st.write(
        "Format your references correctly according to APA, IEEE, or MLA styles without inventing missing details."
    )

    if "references_list" not in st.session_state:
        st.session_state.references_list = []

    with st.form("citation_form"):
        col_c1, col_c2 = st.columns(2)
        with col_c1:
            authors = st.text_input("Author Name(s)", placeholder="e.g., Smith, J. A., & Doe, B.")
            pub_year = st.text_input("Publication Year", placeholder="e.g., 2024")
            title = st.text_input("Article or Book Title", placeholder="e.g., Machine Learning in Structural Biology")
        with col_c2:
            journal = st.text_input("Journal, Publisher, or Conference Name", placeholder="e.g., Journal of Bioinformatics")
            vol_iss_pages = st.text_input("Volume, Issue, and Page Numbers", placeholder="e.g., Vol. 12, Issue 3, pp. 45-58")
            doi_url = st.text_input("DOI or URL", placeholder="e.g., https://doi.org/10.1038/s41586-021-0000")

        submit_citation = st.form_submit_button("Format and Add Reference")

        if submit_citation:
            if not authors.strip() or not title.strip():
                st.warning("Please provide at least Author Name(s) and Article/Book Title.")
            else:
                if not api_key:
                    st.error("Please enter a valid Gemini API Key in the sidebar or secrets.")
                else:
                    client = get_gemini_client(api_key)
                    if not client:
                        st.error("Failed to initialize Gemini client.")
                    else:
                        with st.spinner("Formatting citation..."):
                            try:
                                prompt = f"""
                                You are an expert academic citation formatter. Format the given reference details strictly in the {citation_style} citation style.
                                
                                Reference Details:
                                - Authors: {authors if authors else 'Missing'}
                                - Year: {pub_year if pub_year else 'Missing'}
                                - Title: {title if title else 'Missing'}
                                - Source/Journal: {journal if journal else 'Missing'}
                                - Volume/Issue/Pages: {vol_iss_pages if vol_iss_pages else 'Missing'}
                                - DOI/URL: {doi_url if doi_url else 'Missing'}
                                
                                RULES:
                                1. Do not invent missing authors, publication details, DOIs, or sources.
                                2. Clearly flag any incomplete or missing reference details.
                                3. Provide both the full bibliographic citation AND an appropriate in-text citation format.
                                
                                Output format:
                                Full Citation: [formatted citation]
                                In-Text Citation: [formatted in-text citation]
                                """
                                response = client.models.generate_content(
                                    model="gemini-3.8-flash",
                                    contents=prompt,
                                )
                                formatted_result = response.text
                                st.session_state.references_list.append(formatted_result)
                                st.success("Reference added successfully!")
                            except Exception as e:
                                st.error(f"Error formatting citation: {e}")

    if st.session_state.references_list:
        st.subheader("Formatted References List")
        for idx, ref in enumerate(st.session_state.references_list, 1):
            st.markdown(f"**[{idx}]** {ref}")

        if st.button("Clear All References"):
            st.session_state.references_list = []
            st.rerun()

        refs_combined = "\n\n".join(st.session_state.references_list)
        docx_refs = create_docx("References", refs_combined)
        st.download_button(
            label="Download References as DOCX",
            data=docx_refs,
            file_name="formatted_references.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

# Footer Disclaimer
st.markdown("---")
st.markdown(
    "<div style='text-align: center; color: #6b7280; font-size: 0.85rem;'>"
    "⚠️ <b>Disclaimer:</b> AI-generated academic content must be carefully verified, reviewed, and edited by the student before inclusion in official submissions."
    "</div>",
    unsafe_allow_html=True,
)